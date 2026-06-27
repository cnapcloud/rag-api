"""Unit tests for doc_created_at extraction, upsert payload, meta storage, and reindex."""

from __future__ import annotations

import asyncio
from datetime import datetime, timezone
from unittest.mock import MagicMock, patch

import pytest

DOC_ID = "11111111-1111-1111-1111-111111111111"
STORAGE_KEY = "kb-test/doc.pdf"


# ──────────────────────────────────────────────
# _extract_doc_created_at
# ──────────────────────────────────────────────

class TestExtractDocCreatedAt:
    def _call(self, file_path, suffix, storage_key=STORAGE_KEY):
        from pipeline.ops.parse import _extract_doc_created_at
        return _extract_doc_created_at(file_path, suffix, storage_key)

    def test_pdf_creation_date(self, tmp_path):
        """PDF with CreationDate returns that date as ISO UTC string."""
        from pypdf import PdfWriter

        pdf_path = tmp_path / "test.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        dt = datetime(2023, 5, 15, 10, 30, 0, tzinfo=timezone.utc)
        writer.add_metadata({"/CreationDate": dt.strftime("D:%Y%m%d%H%M%SZ")})
        with open(pdf_path, "wb") as f:
            writer.write(f)

        result = self._call(pdf_path, ".pdf")
        assert "2023-05-15" in result

    def test_pdf_no_creation_date_falls_back_to_s3(self, tmp_path):
        """PDF without CreationDate falls back to S3 LastModified."""
        from pypdf import PdfWriter

        pdf_path = tmp_path / "test.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        with open(pdf_path, "wb") as f:
            writer.write(f)

        with patch("infra.s3.get_object_last_modified_by_key", return_value="2024-01-01T00:00:00+00:00"):
            result = self._call(pdf_path, ".pdf")
        assert result == "2024-01-01T00:00:00+00:00"

    def test_docx_created_date(self, tmp_path):
        """DOCX with core_properties.created returns that date."""
        import docx

        docx_path = tmp_path / "test.docx"
        document = docx.Document()
        document.add_paragraph("Hello")
        document.save(str(docx_path))

        dt = datetime(2022, 3, 10, 8, 0, 0, tzinfo=timezone.utc)
        with patch("docx.Document") as mock_doc_cls:
            mock_doc = MagicMock()
            mock_doc.core_properties.created = dt
            mock_doc_cls.return_value = mock_doc
            result = self._call(docx_path, ".docx")

        assert "2022-03-10" in result

    def test_unsupported_type_falls_back_to_s3(self, tmp_path):
        """Non-PDF/DOCX file types fall back to S3 LastModified."""
        txt_path = tmp_path / "test.txt"
        txt_path.write_text("hello")

        with patch("infra.s3.get_object_last_modified_by_key", return_value="2024-06-01T12:00:00+00:00"):
            result = self._call(txt_path, ".txt")
        assert result == "2024-06-01T12:00:00+00:00"

    def test_s3_fallback_failure_returns_empty(self, tmp_path):
        """If both extraction and S3 fallback fail, return empty string."""
        txt_path = tmp_path / "test.txt"
        txt_path.write_text("hello")

        with patch("infra.s3.get_object_last_modified_by_key", side_effect=Exception("S3 error")):
            result = self._call(txt_path, ".txt")
        assert result == ""

    def test_naive_datetime_coerced_to_utc(self, tmp_path):
        """Naive datetime from PDF metadata is treated as UTC."""
        dt_naive = datetime(2021, 1, 1, 0, 0, 0)

        with patch("pypdf.PdfReader") as mock_reader_cls:
            mock_meta = MagicMock()
            mock_meta.creation_date = dt_naive
            mock_reader_cls.return_value.metadata = mock_meta

            pdf_path = tmp_path / "naive.pdf"
            pdf_path.write_bytes(b"%PDF-1.4")
            result = self._call(pdf_path, ".pdf")

        assert "+00:00" in result or "Z" in result
        assert "2021-01-01" in result


# ──────────────────────────────────────────────
# UpsertResult.doc_created_at + Qdrant payload
# ──────────────────────────────────────────────

class TestUpsertDocCreatedAt:
    def _make_embedded_node(self, doc_created_at: str = ""):
        from llama_index.core.schema import TextNode
        from pipeline.ops.embed import EmbeddedNode

        node = TextNode(text="sample chunk", metadata={"doc_created_at": doc_created_at})
        return EmbeddedNode(
            node=node,
            dense_vector=[0.1] * 4,
            sparse_indices=[0, 1],
            sparse_values=[0.5, 0.5],
        )

    def test_upsert_result_carries_doc_created_at(self, mock_qdrant):
        from pipeline.ops.upsert import upsert

        en = self._make_embedded_node("2023-05-15T10:30:00+00:00")
        with (
            patch("pipeline.ops.upsert.qdrant_infra.get_qdrant_client", return_value=mock_qdrant),
            patch("pipeline.ops.upsert.qdrant_infra.ensure_collection"),
            patch("pipeline.ops.upsert.qdrant_infra.delete_chunks_by_doc_id"),
            patch("pipeline.ops.upsert.qdrant_infra.upsert_chunks"),
        ):
            result = upsert("kb-test", DOC_ID, [en])

        assert result.doc_created_at == "2023-05-15T10:30:00+00:00"

    def test_qdrant_payload_includes_doc_created_at(self, mock_qdrant):
        from pipeline.ops.upsert import upsert

        expected = "2023-05-15T10:30:00+00:00"
        en = self._make_embedded_node(expected)

        captured_points = []

        def capture_upsert(kb_id, points, client=None):
            captured_points.extend(points)

        with (
            patch("pipeline.ops.upsert.qdrant_infra.get_qdrant_client", return_value=mock_qdrant),
            patch("pipeline.ops.upsert.qdrant_infra.ensure_collection"),
            patch("pipeline.ops.upsert.qdrant_infra.delete_chunks_by_doc_id"),
            patch("pipeline.ops.upsert.qdrant_infra.upsert_chunks", side_effect=capture_upsert),
        ):
            upsert("kb-test", DOC_ID, [en])

        assert len(captured_points) == 1
        assert captured_points[0].payload["doc_created_at"] == expected

    def test_qdrant_payload_uses_doc_id_not_doc_key(self, mock_qdrant):
        """Payload contains doc_id field (not doc_key or doc_source)."""
        from pipeline.ops.upsert import upsert

        en = self._make_embedded_node("2023-05-15T10:30:00+00:00")
        captured_points = []

        with (
            patch("pipeline.ops.upsert.qdrant_infra.get_qdrant_client", return_value=mock_qdrant),
            patch("pipeline.ops.upsert.qdrant_infra.ensure_collection"),
            patch("pipeline.ops.upsert.qdrant_infra.delete_chunks_by_doc_id"),
            patch("pipeline.ops.upsert.qdrant_infra.upsert_chunks", side_effect=lambda kb, pts, client=None: captured_points.extend(pts)),
        ):
            upsert("kb-test", DOC_ID, [en])

        payload = captured_points[0].payload
        assert payload.get("doc_id") == DOC_ID
        assert "doc_key" not in payload
        assert "doc_source" not in payload

    def test_empty_embedded_nodes_doc_created_at_is_empty(self, mock_qdrant):
        from pipeline.ops.upsert import upsert

        with (
            patch("pipeline.ops.upsert.qdrant_infra.get_qdrant_client", return_value=mock_qdrant),
            patch("pipeline.ops.upsert.qdrant_infra.ensure_collection"),
            patch("pipeline.ops.upsert.qdrant_infra.delete_chunks_by_doc_id"),
            patch("pipeline.ops.upsert.qdrant_infra.upsert_chunks"),
        ):
            result = upsert("kb-test", DOC_ID, [])

        assert result.doc_created_at == ""


# ──────────────────────────────────────────────
# parse_op stores doc_created_at in Postgres immediately
# ──────────────────────────────────────────────

class TestMetaDocCreatedAt:
    def test_parse_op_saves_doc_created_at(self):
        """parse_op persists doc_created_at to DB right after parse, before dedup."""
        from llama_index.core.schema import Document

        doc = Document(text="hello", metadata={"doc_created_at": "2023-05-15T10:30:00+00:00", "file_name": "x.md"})
        stored: dict = {}

        with patch("infra.postgres.update_doc_fields", side_effect=lambda doc_id, fields: stored.update(fields)), \
             patch("pipeline.ops.parse.parse", return_value=[doc]):
            from dagster import build_op_context
            from defs.ops.ingest_ops import parse_op
            ctx = build_op_context()
            parse_op(ctx, {"doc_id": DOC_ID, "storage_key": "kb/x.md"})

        assert stored.get("doc_created_at") == "2023-05-15T10:30:00+00:00"

    def test_parse_op_skips_save_when_doc_created_at_empty(self):
        """parse_op does not call update_doc_fields when doc_created_at is empty."""
        from llama_index.core.schema import Document

        doc = Document(text="hello", metadata={"doc_created_at": "", "file_name": "x.md"})

        with patch("infra.postgres.update_doc_fields") as mock_udf, \
             patch("pipeline.ops.parse.parse", return_value=[doc]):
            from dagster import build_op_context
            from defs.ops.ingest_ops import parse_op
            ctx = build_op_context()
            parse_op(ctx, {"doc_id": DOC_ID, "storage_key": "kb/x.md"})

        mock_udf.assert_not_called()


# ──────────────────────────────────────────────
# reindex_kb behavior
# ──────────────────────────────────────────────

class TestReindexKb:
    def test_reindex_skips_docs_with_matching_etag(self):
        """Docs with matching S3 ETag and content_version are skipped."""
        docs = [
            {"doc_id": DOC_ID, "kb_id": "kb-test", "storage_key": "kb-test/a.pdf", "content_version": "etag-a"},
        ]

        enqueued = []

        with (
            patch("infra.postgres.list_docs", return_value=docs),
            patch("infra.s3.get_object_meta", return_value=("etag-a", 1024)),
            patch("pipeline.enqueue.enqueue_upload_event", side_effect=lambda doc_id, force=False: enqueued.append(doc_id)),
        ):
            from api.routers.docs import reindex_kb
            result = asyncio.run(reindex_kb(kb_id="kb-test", force=False))

        assert result["skipped"] == 1
        assert result["queued"] == 0
        assert len(enqueued) == 0

    def test_reindex_enqueues_docs_with_changed_etag(self):
        """Docs with different S3 ETag are enqueued."""
        docs = [
            {"doc_id": DOC_ID, "kb_id": "kb-test", "storage_key": "kb-test/a.pdf", "content_version": "old-etag"},
        ]

        enqueued = []

        with (
            patch("infra.postgres.list_docs", return_value=docs),
            patch("infra.s3.get_object_meta", return_value=("new-etag", 1024)),
            patch("pipeline.enqueue.enqueue_upload_event", side_effect=lambda doc_id, force=False: enqueued.append(doc_id)),
        ):
            from api.routers.docs import reindex_kb
            result = asyncio.run(reindex_kb(kb_id="kb-test", force=False))

        assert result["queued"] == 1
        assert result["skipped"] == 0
        assert DOC_ID in enqueued

    def test_reindex_force_enqueues_all(self):
        """force=True enqueues all docs regardless of ETag."""
        docs = [
            {"doc_id": DOC_ID, "kb_id": "kb-test", "storage_key": "kb-test/a.pdf", "content_version": "etag-a"},
        ]

        enqueued = []

        with (
            patch("infra.postgres.list_docs", return_value=docs),
            patch("pipeline.enqueue.enqueue_upload_event", side_effect=lambda doc_id, force=False: enqueued.append(doc_id)),
        ):
            from api.routers.docs import reindex_kb
            result = asyncio.run(reindex_kb(kb_id="kb-test", force=True))

        assert result["queued"] == 1
        assert DOC_ID in enqueued

    def test_reindex_skips_docs_without_storage_key(self):
        """Docs with no storage_key are skipped."""
        docs = [
            {"doc_id": DOC_ID, "kb_id": "kb-test", "storage_key": None, "content_version": None},
        ]

        enqueued = []

        with (
            patch("infra.postgres.list_docs", return_value=docs),
            patch("pipeline.enqueue.enqueue_upload_event", side_effect=lambda doc_id, force=False: enqueued.append(doc_id)),
        ):
            from api.routers.docs import reindex_kb
            result = asyncio.run(reindex_kb(kb_id="kb-test", force=False))

        assert result["skipped"] == 1
        assert result["queued"] == 0
