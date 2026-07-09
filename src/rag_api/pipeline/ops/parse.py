"""parse Op — LlamaIndex SimpleDirectoryReader-based document parsing."""

from __future__ import annotations

import logging
import tempfile
from datetime import datetime, timezone
from pathlib import Path

from llama_index.core import Document, SimpleDirectoryReader
from llama_index.core.readers.base import BaseReader

from rag_api.exceptions import IngestValidationError
from rag_api.infra.s3 import download_by_key

logger = logging.getLogger(__name__)

DOCUMENT_EXTENSIONS: frozenset[str] = frozenset({
    ".pdf", ".md", ".docx", ".txt", ".hwp", ".html", ".htm", ".rst",
})

CODE_EXTENSIONS: frozenset[str] = frozenset({
    ".py", ".ts", ".tsx", ".js", ".jsx",
    ".go", ".java", ".rs",
    ".cpp", ".cc", ".c", ".cs",
    ".rb", ".php", ".swift", ".kt", ".scala", ".sh",
})

CONFIG_DATA_EXTENSIONS: frozenset[str] = frozenset({
    ".yaml", ".yml", ".properties",
})

SUPPORTED_EXTENSIONS = DOCUMENT_EXTENSIONS | CODE_EXTENSIONS | CONFIG_DATA_EXTENSIONS

CODE_LANGUAGE_MAP: dict[str, str] = {
    ".py": "python",
    ".ts": "typescript", ".tsx": "typescript",
    ".js": "javascript", ".jsx": "javascript",
    ".go": "go",
    ".java": "java",
    ".rs": "rust",
    ".cpp": "cpp", ".cc": "cpp",
    ".c": "c",
    ".cs": "c_sharp",
    ".rb": "ruby",
    ".php": "php",
    ".swift": "swift",
    ".kt": "kotlin",
    ".scala": "scala",
    ".sh": "bash",
}


def _extract_doc_created_at(file_path: Path, suffix: str, storage_key: str) -> str:
    """Extract document creation date from file metadata.

    Priority: PDF CreationDate / DOCX core_properties.created -> S3 LastModified fallback.
    Returns an ISO 8601 UTC string, or empty string if all sources fail.
    """
    dt: datetime | None = None

    try:
        if suffix == ".pdf":
            import pypdf

            reader = pypdf.PdfReader(str(file_path))
            if reader.metadata and reader.metadata.creation_date:
                raw = reader.metadata.creation_date
                dt = raw if isinstance(raw, datetime) else None
        elif suffix == ".docx":
            import docx

            doc = docx.Document(str(file_path))
            created = doc.core_properties.created
            if created:
                dt = created if isinstance(created, datetime) else None
    except Exception as e:
        logger.debug("doc_created_at extraction failed (will fallback): file=%s err=%s", file_path.name, e)

    if dt is not None:
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt.astimezone(timezone.utc).isoformat()

    try:
        from rag_api.infra.s3 import get_object_last_modified_by_key

        return get_object_last_modified_by_key(storage_key)
    except Exception as e:
        logger.debug("S3 LastModified fallback failed: storage_key=%s err=%s", storage_key, e)
        return ""


from llama_index.core.readers.base import BaseReader
from llama_index.core.schema import Document


class HTMLCleanReader(BaseReader):
    """HTML reader that extracts main content via trafilatura's density-based
    algorithm instead of a tag-name deny-list, since markup conventions vary
    too much across sites for deny-listing to reliably strip boilerplate.

    favor_precision (settings.ingestion.html_favor_precision) excludes ambiguous
    blocks (sidebar-or-content boundary cases) so near-duplicate pages hash
    consistently in dedup stage 1 (SimHash). output_format="markdown" preserves
    heading/list structure for downstream chunking, instead of flattening
    everything with get_text().
    """

    def load_data(self, file: Path, extra_info: dict | None = None) -> list[Document]:
        import trafilatura

        from rag_api.config.settings import get_settings

        with open(file, encoding="utf-8") as f:
            html = f.read()

        favor_precision = get_settings().ingestion.html_favor_precision
        text = (
            trafilatura.extract(
                html,
                favor_precision=favor_precision,
                output_format="markdown",
                include_tables=True,
            )
            or ""
        )

        metadata: dict = {"file_path": str(file)}
        metadata.update(extra_info or {})

        return [Document(text=text, metadata=metadata)]


class PyMuPDFReader(BaseReader):
    """PDF reader using PyMuPDF (fitz) instead of pypdf.

    Keynote/Pages PDF exports often embed subset CID fonts without a valid
    ToUnicode CMap, so pypdf's glyph-to-unicode mapping (extract_doc.py's
    PDFReader) returns garbled or empty text for CJK content while ASCII
    text extracts fine. PyMuPDF resolves glyphs via the font's own cmap
    table and handles this case correctly.
    """

    def load_data(self, file: Path, extra_info: dict | None = None) -> list[Document]:
        import fitz

        docs = []
        with fitz.open(str(file)) as pdf:
            for page_num, page in enumerate(pdf, start=1):
                metadata = {"page_label": str(page_num), "file_name": Path(file).name}
                if extra_info:
                    metadata.update(extra_info)
                docs.append(Document(text=page.get_text(), metadata=metadata))
        return docs


def _get_file_extractor() -> dict:
    try:
        from llama_index.readers.file import DocxReader, FlatReader, MarkdownReader
        from llama_index.readers.hwp import HWPReader
    except ImportError:
        logger.warning("llama-index-readers-file not installed, using default reader")
        return {}

    extractors: dict = {
        ".pdf": PyMuPDFReader(),
        ".md": MarkdownReader(),
        ".docx": DocxReader(),
        ".txt": FlatReader(),
        ".hwp": HWPReader(),
        ".html": HTMLCleanReader(),
        ".htm": HTMLCleanReader(),
        ".rst": FlatReader(),
    }
    for ext in CODE_EXTENSIONS:
        extractors[ext] = FlatReader()
    return extractors


def parse(doc_id: str, storage_key: str, local_path: Path | None = None) -> list[Document]:
    """Download a file from S3 and return a list of LlamaIndex Documents.

    Args:
        doc_id: UUID of the document row (used to tag metadata).
        storage_key: Full S3 object path (e.g. 'kb-01/report.pdf').
        local_path: Pre-downloaded local file (for tests / CLI use).
    """
    suffix = Path(storage_key).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise IngestValidationError(f"Unsupported file format: {suffix}")

    with tempfile.TemporaryDirectory() as tmpdir:
        if local_path is None:
            dest = Path(tmpdir) / Path(storage_key).name
            download_by_key(storage_key, dest)
            file_path = dest
        else:
            file_path = local_path

        extractor = _get_file_extractor()
        reader = SimpleDirectoryReader(
            input_files=[str(file_path)],
            file_extractor=extractor if extractor else None,
        )
        documents = reader.load_data()

        doc_created_at = _extract_doc_created_at(file_path, suffix, storage_key)

        for doc in documents:
            doc.metadata.update(
                {
                    "doc_id": doc_id,
                    "doc_type": suffix.lstrip("."),
                    "doc_created_at": doc_created_at,
                }
            )

        logger.info(
            "Parsed: doc_id=%s storage_key=%s documents=%d doc_created_at=%s",
            doc_id, storage_key, len(documents), doc_created_at or "n/a",
        )
        return documents


def parse_local(file_path: Path, doc_id: str = "local", storage_key: str | None = None) -> list[Document]:
    """Parse a local file directly (CLI / test use)."""
    key = storage_key or file_path.name
    return parse(doc_id=doc_id, storage_key=key, local_path=file_path)
